"""
Report Service — Milestone 4 Step 5
Generates PDF and Excel reports using real internal SocialPilot metrics.
All metrics match AnalyticsService calculations for consistency.
Strictly team/workspace isolated with authorization enforcement.
Zero fake engagement data. Zero secret token leakage.
"""

import io
import logging
from datetime import datetime, timezone, timedelta
from typing import Dict, Any, List, Optional

from sqlalchemy.orm import Session
from motor.motor_asyncio import AsyncIOMotorDatabase

from app.models.user import User
from app.services.team_service import TeamService
from app.services.analytics_service import (
    get_overview as get_analytics_overview,
    get_campaign_performance as get_campaign_analytics,
)
from app.db.mongo import ensure_active_mongo_db

# PDF generation (ReportLab)
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# Excel generation (OpenPyXL)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ReportService:
    @staticmethod
    async def gather_report_data(
        db: Session,
        mongo_db: AsyncIOMotorDatabase,
        user: User,
        team_id: int,
        days: int = 30,
    ) -> Dict[str, Any]:
        """
        Gathers verified team metadata, analytics overview, campaign performance,
        and post list within the specified date range.
        Enforces team membership authorization via TeamService.
        """
        # 1. Authorize team access & get team entity
        team = TeamService.get_team_by_id(db, team_id, user)

        # 2. Get analytics overview and campaign performance
        overview = await get_analytics_overview(db, mongo_db, user, team_id, days)
        campaign_perf = await get_campaign_analytics(db, mongo_db, user, team_id)

        # 3. Query individual posts for period (for detailed Excel Posts sheet)
        mongo_db = ensure_active_mongo_db(mongo_db)
        posts_coll = mongo_db["posts"]

        now = datetime.now(timezone.utc)
        period_start = now - timedelta(days=days)

        posts_filter = {
            "team_id": team_id,
            "created_at": {"$gte": period_start},
        }

        posts_docs = (
            await posts_coll.find(posts_filter)
            .sort("created_at", -1)
            .to_list(length=1000)
        )

        formatted_posts = []
        for p in posts_docs:
            created_at_val = p.get("created_at")
            created_at_str = (
                created_at_val.isoformat()
                if isinstance(created_at_val, datetime)
                else str(created_at_val or "")
            )

            scheduled_at_val = p.get("scheduled_at")
            scheduled_at_str = (
                scheduled_at_val.isoformat()
                if isinstance(scheduled_at_val, datetime)
                else str(scheduled_at_val or "")
            )

            published_at_val = p.get("published_at")
            published_at_str = (
                published_at_val.isoformat()
                if isinstance(published_at_val, datetime)
                else str(published_at_val or "")
            )

            formatted_posts.append({
                "post_id": str(p.get("_id", "")),
                "content": p.get("content") or p.get("base_content") or "",
                "status": p.get("status", "draft"),
                "target_platforms": ", ".join(p.get("target_platforms", [])),
                "created_at": created_at_str,
                "scheduled_at": scheduled_at_str,
                "published_at": published_at_str,
                "campaign_id": str(p.get("campaign_id", "")) if p.get("campaign_id") else "-",
            })

        return {
            "team_id": team.id,
            "team_name": team.name,
            "generated_at": now.strftime("%Y-%m-%d %H:%M:%S UTC"),
            "period_days": days,
            "period_start": period_start.strftime("%Y-%m-%d"),
            "period_end": now.strftime("%Y-%m-%d"),
            "overview": overview,
            "campaign_performance": campaign_perf.get("campaigns", []),
            "posts": formatted_posts,
        }

    @staticmethod
    def generate_pdf_report(report_data: Dict[str, Any]) -> bytes:
        """
        Generates a clean, professional PDF report using ReportLab.
        Returns raw bytes of the generated PDF.
        """
        buffer = io.BytesIO()

        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()

        # Custom Palette
        PRIMARY_COLOR = colors.HexColor("#15803D")   # Green accent
        TEXT_MAIN = colors.HexColor("#0F172A")       # Dark slate
        TEXT_MUTED = colors.HexColor("#64748B")      # Muted gray
        BG_CARD = colors.HexColor("#F8FAFC")         # Light surface
        BORDER_COLOR = colors.HexColor("#E2E8F0")    # Soft border

        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=PRIMARY_COLOR,
            spaceAfter=4,
        )

        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=TEXT_MUTED,
            spaceAfter=12,
        )

        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=17,
            textColor=TEXT_MAIN,
            spaceBefore=12,
            spaceAfter=6,
        )

        body_style = ParagraphStyle(
            'ReportBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=TEXT_MAIN,
        )

        table_header_style = ParagraphStyle(
            'TableHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=9,
            leading=11,
            textColor=colors.white,
        )

        table_cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11,
            textColor=TEXT_MAIN,
        )

        story = []

        # 1. Header Banner
        story.append(Paragraph("SocialPilot Workspace Report", title_style))
        sub_text = (
            f"Workspace: <b>{report_data['team_name']}</b> | "
            f"Report Period: <b>Last {report_data['period_days']} Days</b> "
            f"({report_data['period_start']} to {report_data['period_end']}) | "
            f"Generated: {report_data['generated_at']}"
        )
        story.append(Paragraph(sub_text, subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY_COLOR, spaceAfter=14))

        # 2. Executive Summary Metrics
        story.append(Paragraph("1. Executive Summary", section_heading))

        posts_meta = report_data['overview']['posts']
        summary_data = [
            [
                Paragraph("<b>Total Posts</b>", table_header_style),
                Paragraph("<b>Published</b>", table_header_style),
                Paragraph("<b>Scheduled</b>", table_header_style),
                Paragraph("<b>Draft</b>", table_header_style),
                Paragraph("<b>Failed</b>", table_header_style),
                Paragraph("<b>Success Rate</b>", table_header_style),
            ],
            [
                Paragraph(str(posts_meta['period_total']), table_cell_style),
                Paragraph(str(posts_meta['published']), table_cell_style),
                Paragraph(str(posts_meta['scheduled']), table_cell_style),
                Paragraph(str(posts_meta['draft']), table_cell_style),
                Paragraph(str(posts_meta['failed']), table_cell_style),
                Paragraph(f"{posts_meta['success_rate']}%", table_cell_style),
            ],
        ]

        summary_table = Table(summary_data, colWidths=[90, 90, 90, 90, 90, 90])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY_COLOR),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BACKGROUND', (0, 1), (-1, 1), BG_CARD),
            ('GRID', (0, 0), (-1, -1), 0.5, BORDER_COLOR),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 14))

        # 3. Platform Distribution
        story.append(Paragraph("2. Platform Post Distribution", section_heading))
        platforms = report_data['overview'].get('platform_breakdown', [])

        if platforms:
            platform_rows = [
                [
                    Paragraph("<b>Platform</b>", table_header_style),
                    Paragraph("<b>Post Count</b>", table_header_style),
                    Paragraph("<b>Share of Total</b>", table_header_style),
                ]
            ]
            total_p_posts = max(1, posts_meta['period_total'])
            for item in platforms:
                p_name = item['platform'].capitalize()
                cnt = item['count']
                share = round(cnt / total_p_posts * 100, 1)
                platform_rows.append([
                    Paragraph(p_name, table_cell_style),
                    Paragraph(str(cnt), table_cell_style),
                    Paragraph(f"{share}%", table_cell_style),
                ])
            p_table = Table(platform_rows, colWidths=[180, 180, 180])
            p_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), PRIMARY_COLOR),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('BACKGROUND', (0, 1), (-1, -1), BG_CARD),
                ('GRID', (0, 0), (-1, -1), 0.5, BORDER_COLOR),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(p_table)
        else:
            story.append(Paragraph("<i>No platform post activity recorded in this period.</i>", body_style))

        story.append(Spacer(1, 14))

        # 4. Campaign Summary
        story.append(Paragraph("3. Campaign Performance Summary", section_heading))
        c_perf = report_data.get('campaign_performance', [])
        if c_perf:
            c_rows = [
                [
                    Paragraph("<b>Campaign Name</b>", table_header_style),
                    Paragraph("<b>Status</b>", table_header_style),
                    Paragraph("<b>Total Posts</b>", table_header_style),
                    Paragraph("<b>Published</b>", table_header_style),
                    Paragraph("<b>Scheduled</b>", table_header_style),
                    Paragraph("<b>Publish Rate</b>", table_header_style),
                ]
            ]
            for c in c_perf:
                c_rows.append([
                    Paragraph(c['name'], table_cell_style),
                    Paragraph(c['status'].capitalize(), table_cell_style),
                    Paragraph(str(c['total_posts']), table_cell_style),
                    Paragraph(str(c['published_posts']), table_cell_style),
                    Paragraph(str(c['scheduled_posts']), table_cell_style),
                    Paragraph(f"{c['publish_rate']}%", table_cell_style),
                ])
            c_table = Table(c_rows, colWidths=[160, 70, 70, 70, 70, 100])
            c_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), PRIMARY_COLOR),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('BACKGROUND', (0, 1), (-1, -1), BG_CARD),
                ('GRID', (0, 0), (-1, -1), 0.5, BORDER_COLOR),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(c_table)
        else:
            story.append(Paragraph("<i>No campaigns found for this workspace.</i>", body_style))

        story.append(Spacer(1, 14))

        # 5. Connected Social Accounts
        story.append(Paragraph("4. Connected Social Accounts Health", section_heading))
        acct_meta = report_data['overview'].get('accounts', {})
        acct_rows = [
            [
                Paragraph("<b>Metric</b>", table_header_style),
                Paragraph("<b>Value</b>", table_header_style),
            ],
            [Paragraph("Total Connected Accounts", table_cell_style), Paragraph(str(acct_meta.get('connected', 0)), table_cell_style)],
            [Paragraph("Expired / Re-auth Required", table_cell_style), Paragraph(str(acct_meta.get('expired', 0)), table_cell_style)],
            [Paragraph("Error Accounts", table_cell_style), Paragraph(str(acct_meta.get('error', 0)), table_cell_style)],
            [Paragraph("Total Workspace Accounts", table_cell_style), Paragraph(str(acct_meta.get('total', 0)), table_cell_style)],
        ]
        acct_table = Table(acct_rows, colWidths=[270, 270])
        acct_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), PRIMARY_COLOR),
            ('BACKGROUND', (0, 1), (-1, -1), BG_CARD),
            ('GRID', (0, 0), (-1, -1), 0.5, BORDER_COLOR),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(acct_table)

        story.append(Spacer(1, 16))

        # 6. External Analytics Disclosure Notice
        notice_box_style = ParagraphStyle(
            'NoticeBox',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#92400E"), # Amber-800
        )
        notice_text = (
            "<b>Note on External Platform Engagement:</b> "
            "External engagement metrics (likes, reach, impressions, comments, CTR) are currently "
            "unavailable without active provider analytics OAuth scopes. Connect social accounts with "
            "read_insights permissions to display engagement metrics."
        )

        notice_table = Table(
            [[Paragraph(notice_text, notice_box_style)]],
            colWidths=[540],
        )
        notice_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FEF3C7")), # Amber-100
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#F59E0B")),      # Amber-500
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
        ]))
        story.append(notice_table)

        # Build document
        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    @staticmethod
    def generate_excel_report(report_data: Dict[str, Any]) -> bytes:
        """
        Generates a clean, formatted .xlsx Excel workbook using OpenPyXL.
        Sheets: Summary, Posts, Platform Breakdown, Campaigns, Social Accounts, Daily Trend.
        Returns raw bytes of the generated workbook.
        """
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Styles
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='15803D', end_color='15803D', fill_type='solid')
        title_font = Font(name='Calibri', size=14, bold=True, color='0F172A')
        subtitle_font = Font(name='Calibri', size=10, italic=True, color='64748B')
        bold_font = Font(name='Calibri', size=11, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0'),
        )
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')

        # -------------------------------------------------------------
        # SHEET 1: Summary
        # -------------------------------------------------------------
        ws_sum = wb.create_sheet(title="Summary")
        ws_sum.views.sheetView[0].showGridLines = True

        ws_sum.append(["SocialPilot Workspace Analytics Summary"])
        ws_sum.cell(row=1, column=1).font = title_font

        ws_sum.append([f"Workspace: {report_data['team_name']} | Period: Last {report_data['period_days']} Days ({report_data['period_start']} to {report_data['period_end']})"])
        ws_sum.cell(row=2, column=1).font = subtitle_font

        ws_sum.append([f"Generated At: {report_data['generated_at']}"])
        ws_sum.cell(row=3, column=1).font = subtitle_font
        ws_sum.append([])

        # Post Metrics Table
        ws_sum.append(["Post Publishing Overview"])
        ws_sum.cell(row=5, column=1).font = bold_font

        ws_sum.append(["Metric", "Count / Value"])
        for col in range(1, 3):
            cell = ws_sum.cell(row=6, column=col)
            cell.font = header_font
            cell.fill = header_fill

        posts_meta = report_data['overview']['posts']
        metrics_rows = [
            ("Period Total Posts", posts_meta['period_total']),
            ("Published Posts", posts_meta['published']),
            ("Scheduled Posts", posts_meta['scheduled']),
            ("Draft Posts", posts_meta['draft']),
            ("Failed Posts", posts_meta['failed']),
            ("Publish Success Rate", f"{posts_meta['success_rate']}%"),
            ("Lifetime Total Posts", posts_meta['lifetime_total']),
            ("Lifetime Published Posts", posts_meta['lifetime_published']),
        ]
        for name, val in metrics_rows:
            ws_sum.append([name, val])

        ws_sum.append([])
        ws_sum.append(["External Engagement Analytics Notice"])
        ws_sum.cell(row=ws_sum.max_row, column=1).font = bold_font
        ws_sum.append([report_data['overview']['engagement']['reason']])

        # -------------------------------------------------------------
        # SHEET 2: Posts
        # -------------------------------------------------------------
        ws_posts = wb.create_sheet(title="Posts")
        ws_posts.views.sheetView[0].showGridLines = True
        post_headers = ["Post ID", "Content", "Status", "Target Platforms", "Created At", "Scheduled At", "Published At", "Campaign ID"]
        ws_posts.append(post_headers)
        for col in range(1, len(post_headers) + 1):
            cell = ws_posts.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        ws_posts.freeze_panes = 'A2'

        for p in report_data.get('posts', []):
            ws_posts.append([
                p['post_id'],
                p['content'],
                p['status'],
                p['target_platforms'],
                p['created_at'],
                p['scheduled_at'],
                p['published_at'],
                p['campaign_id'],
            ])

        # -------------------------------------------------------------
        # SHEET 3: Platform Breakdown
        # -------------------------------------------------------------
        ws_plat = wb.create_sheet(title="Platform Breakdown")
        ws_plat.views.sheetView[0].showGridLines = True
        plat_headers = ["Platform", "Post Count", "Share of Total"]
        ws_plat.append(plat_headers)
        for col in range(1, len(plat_headers) + 1):
            cell = ws_plat.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        ws_plat.freeze_panes = 'A2'

        total_posts_num = max(1, posts_meta['period_total'])
        for item in report_data['overview'].get('platform_breakdown', []):
            p_name = item['platform'].capitalize()
            cnt = item['count']
            share = round(cnt / total_posts_num * 100, 1)
            ws_plat.append([p_name, cnt, f"{share}%"])

        # -------------------------------------------------------------
        # SHEET 4: Campaigns
        # -------------------------------------------------------------
        ws_camp = wb.create_sheet(title="Campaigns")
        ws_camp.views.sheetView[0].showGridLines = True
        camp_headers = ["Campaign ID", "Campaign Name", "Status", "Target Platforms", "Start Date", "End Date", "Total Posts", "Published", "Scheduled", "Failed", "Publish Rate"]
        ws_camp.append(camp_headers)
        for col in range(1, len(camp_headers) + 1):
            cell = ws_camp.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        ws_camp.freeze_panes = 'A2'

        for c in report_data.get('campaign_performance', []):
            ws_camp.append([
                c['campaign_id'],
                c['name'],
                c['status'],
                ", ".join(c.get('target_platforms', [])),
                c['start_date'] or "-",
                c['end_date'] or "-",
                c['total_posts'],
                c['published_posts'],
                c['scheduled_posts'],
                c['failed_posts'],
                f"{c['publish_rate']}%",
            ])

        # -------------------------------------------------------------
        # SHEET 5: Social Accounts
        # -------------------------------------------------------------
        ws_acct = wb.create_sheet(title="Social Accounts")
        ws_acct.views.sheetView[0].showGridLines = True
        acct_headers = ["Metric", "Count"]
        ws_acct.append(acct_headers)
        for col in range(1, len(acct_headers) + 1):
            cell = ws_acct.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

        acct_info = report_data['overview'].get('accounts', {})
        ws_acct.append(["Connected Accounts", acct_info.get('connected', 0)])
        ws_acct.append(["Expired Accounts", acct_info.get('expired', 0)])
        ws_acct.append(["Error Accounts", acct_info.get('error', 0)])
        ws_acct.append(["Total Workspace Accounts", acct_info.get('total', 0)])
        ws_acct.append([])
        ws_acct.append(["Connected Accounts by Platform", "Count"])
        ws_acct.cell(row=ws_acct.max_row, column=1).font = bold_font
        for item in acct_info.get('by_platform', []):
            ws_acct.append([item['platform'].capitalize(), item['count']])

        # -------------------------------------------------------------
        # SHEET 6: Daily Trend
        # -------------------------------------------------------------
        ws_trend = wb.create_sheet(title="Daily Trend")
        ws_trend.views.sheetView[0].showGridLines = True
        trend_headers = ["Date", "Published Posts Count"]
        ws_trend.append(trend_headers)
        for col in range(1, len(trend_headers) + 1):
            cell = ws_trend.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        ws_trend.freeze_panes = 'A2'

        for item in report_data['overview'].get('daily_published_trend', []):
            ws_trend.append([item['date'], item['published']])

        # -------------------------------------------------------------
        # Auto-adjust column widths across all sheets
        # -------------------------------------------------------------
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        buffer = io.BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()
        return excel_bytes
